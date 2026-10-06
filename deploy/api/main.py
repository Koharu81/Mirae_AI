import os, re, json, time, secrets, hashlib, hmac, smtplib, socket, ipaddress, asyncio, base64
from email.message import EmailMessage
from typing import Any
from datetime import datetime, timedelta, timezone, date
from urllib.parse import urlparse
import httpx
from fastapi import FastAPI, Header, HTTPException, Response, Request, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, RedirectResponse
from pydantic import BaseModel, Field, HttpUrl
import psycopg
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from io import BytesIO
import zipfile

APP_VERSION="6.6.0"
app=FastAPI(title="Mirae AI API",version=APP_VERSION,openapi_url=None,docs_url=None,redoc_url=None)
_TRAFFIC={}
def traffic_key(request:Request): return request.headers.get("cf-connecting-ip") or (request.client.host if request.client else "unknown")
def traffic_snapshot():
    now=time.time(); cutoff=now-900; out=[]
    for ip,items in list(_TRAFFIC.items()):
        fresh=[t for t in items if t>=cutoff]
        if fresh:_TRAFFIC[ip]=fresh;out.append((ip,fresh))
        else:_TRAFFIC.pop(ip,None)
    return out
def traffic_record(request:Request):
    ip=traffic_key(request); now=time.time(); arr=_TRAFFIC.setdefault(ip,[]); arr.append(now)
    _TRAFFIC[ip]=[t for t in arr if t>=now-900]
    return ip,len(_TRAFFIC[ip])
ADMIN_EMAIL="admin@koharu.live"
def request_ip(request:Request): return (request.headers.get("cf-connecting-ip") or (request.client.host if request.client else "") or "").strip()
def request_location(request:Request): return {"city":request.headers.get("cf-ipcity","").strip(),"region":request.headers.get("cf-region","").strip(),"country":request.headers.get("cf-ipcountry","").strip()}
def is_admin(request:Request):
    u=session_user(request); return bool(u and str(u.get("email","")).lower()==ADMIN_EMAIL)
def require_admin(request:Request):
    if not is_admin(request): raise HTTPException(403,"관리자 권한이 필요합니다.")
    return session_user(request)
@app.middleware("http")
async def access_audit(request:Request,call_next):
    started=time.perf_counter(); response=None
    ip,burst=traffic_record(request)
    try:
        response=await call_next(request); return response
    finally:
        try:
            path=request.url.path; status=response.status_code if response else 500
            if path not in {"/health","/openapi.json","/docs","/redoc"}:
                u=session_user(request); loc=request_location(request); risk="normal"
                recent=[t for t in _TRAFFIC.get(ip,[]) if t>=time.time()-60]
                if len(recent)>=120: risk="traffic_burst"
                elif status in {401,403}: risk="auth_failure" if path.startswith("/auth/") else "permission_denied"
                elif status==429: risk="rate_limit"
                elif status>=500: risk="server_error"
                elif status>=400: risk="bad_request"
                with db() as c:
                    c.execute("""INSERT INTO mirae_access_logs
                        (user_id,email,nickname,ip,country,region,city,path,method,status,risk_category,risk_detail,user_agent,latency_ms)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                        [u["id"] if u else None,u["email"] if u else "",u["name"] if u else "",request_ip(request),
                         loc["country"],loc["region"],loc["city"],path,request.method,status,risk,
                         ("" if risk=="normal" else "HTTP "+str(status)),request.headers.get("user-agent","")[:1000],
                         int((time.perf_counter()-started)*1000)]); c.commit()
        except Exception: pass
app.add_middleware(CORSMiddleware,allow_origins=["https://gpt-phi-cyan.vercel.app","https://mirae.koharu.live"],allow_credentials=True,allow_methods=["*"],allow_headers=["*"])

@app.get("/docs",include_in_schema=False)
async def custom_docs():
    return RedirectResponse("https://mirae.koharu.live/api-docs.html",status_code=307)

DATABASE_URL=os.getenv("DATABASE_URL","")
HF_TOKEN=os.getenv("HF_TOKEN","")
HF_MODEL_RAW=os.getenv("HF_MODEL","").strip()
UNSUPPORTED_MODELS={"Qwen/Qwen2.5-7B-Instruct","Qwen/Qwen2.5-7B-Instruct:fastest","Qwen/Qwen2.5-7B-Instruct:auto"}
HF_MODEL=HF_MODEL_RAW if HF_MODEL_RAW and HF_MODEL_RAW not in UNSUPPORTED_MODELS else "openai/gpt-oss-120b:groq"
MODEL_API_URL=os.getenv("MODEL_API_URL","").strip()
MODEL_API_KEY=os.getenv("MODEL_API_KEY","").strip()
MODEL_NAME=os.getenv("MODEL_NAME","Mirae-Qwen2.5-1.5B-Instruct").strip() or "Mirae-Qwen2.5-1.5B-Instruct"
NEWS_RSS="https://news.google.com/rss/search"
RESEND_API_KEY=os.getenv("RESEND_API_KEY","")
RESEND_FROM=os.getenv("RESEND_FROM","admin@koharu.live")
RESEND_URL="https://api.resend.com/emails"
CLOUDFLARE_ACCOUNT_ID=os.getenv("CLOUDFLARE_ACCOUNT_ID","").strip()
CLOUDFLARE_API_TOKEN=os.getenv("CLOUDFLARE_API_TOKEN","").strip()
CLOUDFLARE_IMAGE_MODEL=os.getenv("CLOUDFLARE_IMAGE_MODEL","@cf/stabilityai/stable-diffusion-xl-base-1.0").strip()
POLLINATIONS_API_KEY=os.getenv("POLLINATIONS_API_KEY","").strip()
POLLINATIONS_VISION_MODEL=os.getenv("POLLINATIONS_VISION_MODEL","google/gemini-3-flash-preview").strip()
SESSION_DAYS=30

WEB_EXPLICIT=re.compile(r"(웹\s*검색|인터넷(?:에서)?|검색(?:해|해줘|해봐|해서|하고|결과)?|찾아(?:줘|봐|서|서 알려)|공식\s*(?:사이트|자료|문서|페이지)|링크\s*(?:찾|알려)|자료\s*(?:찾|검색))",re.I)
WEB_FRESH=re.compile(r"(최신|현재|지금|최근|실시간|오늘|어제|내일|이번\s*(?:주|달)|업데이트|속보|새로\s*나온)",re.I)
WEB_CONTEXT=re.compile(r"(뉴스|소식|정보|날씨|가격|환율|주가|시세|일정|출시|버전|패치|업데이트|사건|공지|공식|순위|경기|결과|상태|영업|운영시간)",re.I)
GREETING_ONLY=re.compile(r"^\s*(안녕(?:하세요)?|하이|ㅎㅇ|hello|hi|hey|반가워|좋은\s*(?:아침|저녁)|잘\s*지내)\s*[!?.~]*\s*$",re.I)

def db():
    if not DATABASE_URL: raise HTTPException(503,"DATABASE_URL is not configured.")
    return psycopg.connect(DATABASE_URL,row_factory=dict_row)

def digest(v:str)->str: return hashlib.sha256(v.encode()).hexdigest()
def phash(v:str)->str:
    salt=secrets.token_bytes(16)
    return "scrypt$"+salt.hex()+"$"+hashlib.scrypt(v.encode(),salt=salt,n=16384,r=8,p=1).hex()
def pok(v:str,s:str)->bool:
    try:
        _,salt,h=s.split("$",2)
        return hmac.compare_digest(hashlib.scrypt(v.encode(),salt=bytes.fromhex(salt),n=16384,r=8,p=1).hex(),h)
    except Exception:return False

def session_user(request:Request):
    t=request.cookies.get("mirae_session")
    if not t:return None
    with db() as c:return c.execute("SELECT u.* FROM mirae_sessions s JOIN mirae_users u ON u.id=s.user_id WHERE s.token_hash=%s AND s.expires_at>now()",[digest(t)]).fetchone()

class Signup(BaseModel):
    email:str; password:str=Field(min_length=8,max_length=200); name:str=Field(min_length=1,max_length=40)
class Login(BaseModel): email:str; password:str
class Verify(BaseModel): email:str; code:str=Field(min_length=6,max_length=6)
class ChatMessage(BaseModel): role:str; content:str
class ChatAttachment(BaseModel):
    name:str=Field(min_length=1,max_length=180)
    type:str=Field("",max_length=180)
    size:int=Field(0,ge=0,le=25_000_000)
    text:str=Field("",max_length=20_000)
class ChatRequest(BaseModel):
    message:str=Field(min_length=1,max_length=12000); history:list[ChatMessage]=Field(default_factory=list)
    personality:str="balanced"; instructions:str=""; web_search:bool=True
    temperature:float=Field(.7,ge=.2,le=1.2); max_tokens:int=Field(2200,ge=64,le=3200)
    conversation_id:str|None=None
    attachments:list[ChatAttachment]=Field(default_factory=list,max_length=5)
class Settings(BaseModel):
    theme:str="light"; personality:str="balanced"; instructions:str=""; web_search:bool=True
    temperature:float=Field(.7,ge=.2,le=1.2)
class Profile(BaseModel):
    name:str=Field(min_length=1,max_length=40); bio:str=Field("",max_length=500)
    birth_date:date|None=None; avatar_url:str=Field("",max_length=1000)
class KeyCreate(BaseModel): name:str=Field(min_length=2,max_length=80)
class SkillCreate(BaseModel):
    name:str=Field(min_length=1,max_length=60); description:str=Field("",max_length=500)
    url:HttpUrl; method:str=Field("POST",pattern=r"^(GET|POST|PUT|PATCH|DELETE)$")
    headers:str=Field("",max_length=8000); body:str=Field("",max_length=20000)
class SkillRun(BaseModel): params:dict[str,Any]=Field(default_factory=dict)
class ConversationRename(BaseModel): title:str=Field(min_length=1,max_length=120)
class ConversationFolder(BaseModel): name:str=Field(min_length=1,max_length=60)
class ConversationFolderAssign(BaseModel): folder_id:int|None=None
class ConversationFavorite(BaseModel): favorite:bool
class WebhookCreate(BaseModel): name:str=Field(min_length=1,max_length=80); url:str; events:list[str]=Field(default_factory=lambda:["api.request"])
class WebhookState(BaseModel): active:bool
class MemoryCreate(BaseModel): content:str=Field(min_length=1,max_length=1000)
class SkillPromptCreate(BaseModel): prompt:str=Field(min_length=10,max_length=4000)
class PluginCreate(BaseModel):
    name:str=Field(min_length=1,max_length=80); description:str=Field("",max_length=500)
    url:HttpUrl; method:str=Field("POST",pattern=r"^(GET|POST)$")
    permissions:list[str]=Field(default_factory=list,max_length=8)
class PluginState(BaseModel): active:bool
class PluginInvoke(BaseModel): input:dict[str,Any]=Field(default_factory=dict)

def lang(t:str):
    ko=len(re.findall(r"[가-힣]",t));ja=len(re.findall(r"[ぁ-ゖァ-ヺ]",t))
    if ko>=2:return "ko"
    if ja>=2 and ko==0:return "ja"
    return "en"

def wants_web(t:str)->bool:
    text=str(t or "").strip()
    if not text or GREETING_ONLY.fullmatch(text): return False
    return bool(WEB_EXPLICIT.search(text) or WEB_FRESH.search(text) or WEB_CONTEXT.search(text))

def clean_query(t:str)->str:
    t=re.sub(r"(검색해줘|검색해|찾아줘|찾아봐|찾아서|알려줘|알려 줘|정리해줘|정리해 줘|알려|찾아|검색|조회해줘|조회해|확인해줘|확인해|최신|현재|지금|최근|실시간|오늘|어제|내일|이번\s*(?:주|달)|소식|뉴스|정보)"," ",t,flags=re.I)
    q=re.sub(r"\s+"," ",t).strip()
    return q[:180] or "주요 뉴스"

def relevance(q:str,x:dict)->float:
    hay=(x["title"]+" "+x["snippet"]).lower()
    toks=re.findall(r"[가-힣]{2,}|[a-z0-9][a-z0-9._+-]*",q.lower())
    score=sum((2 if len(k)>=3 else 1) for k in toks if k in hay)
    return score

async def search_web(t:str):
    q=clean_query(t)
    async with httpx.AsyncClient(timeout=8,follow_redirects=True) as x:
        r=await x.get(NEWS_RSS,params={"q":q,"hl":"ko","gl":"KR","ceid":"KR:ko"},headers={"User-Agent":"MiraeAI/5.0"})
        r.raise_for_status()
    import xml.etree.ElementTree as ET
    root=ET.fromstring(r.text); out=[]
    seen=set()
    for i in root.findall(".//item"):
        title=" ".join((i.findtext("title") or "").split())
        link=(i.findtext("link") or "").strip(); pub=(i.findtext("pubDate") or "").strip()
        if not title and not link:continue
        if link and link in seen:continue
        if link:seen.add(link)
        desc=re.sub(r"<[^>]+>"," ",i.findtext("description") or "")
        out.append({"title":title,"url":link,"published":pub,"snippet":" ".join(desc.split())[:700]})
    out.sort(key=lambda x:relevance(q,x),reverse=True)
    scored=[x for x in out if relevance(q,x)>0]
    return scored[:8]

def system_prompt(req,language,sources,skills,memories=None,profile_data=None,attachments=None):
    rule={"ko":"한국어로 자연스럽게 답하세요. 사용자가 요청하지 않는 한 다른 언어를 섞지 마세요.","ja":"自然な日本語で答えてください。","en":"Answer in natural English unless the user requests another language."}[language]
    styles={"balanced":"균형 잡힌 말투로 정확성과 자연스러움을 함께 유지하세요.","friendly":"친근하고 편안한 말투를 사용하되 과한 유행어와 이모지는 피하세요.","professional":"전문적이고 구조적인 말투로 답하세요. 필요한 용어는 정확하게 사용하세요.","concise":"핵심부터 짧고 직접적으로 답하세요. 불필요한 반복을 줄이세요.","creative":"아이디어를 적극적으로 확장하되 사실과 창작을 구분하세요."}
    src=""
    if sources:
        src="\n웹 검색 결과:\n"+"\n".join(f"- {x['title']} | {x['published']} | {x['url']} | {x['snippet']}" for x in sources)
    sk="\n사용 가능한 스킬: "+", ".join(f"/skill {x['name']} {{...}}" for x in skills) if skills else ""
    mm=memories or []
    mem="\n기억된 사용자 정보(대화에 도움이 될 때만 사용):\n"+"\n".join(f"- {m['content']}" for m in mm) if mm else ""
    p=profile_data or {}
    profile_text=""
    if p.get("name") or p.get("bio") or p.get("birth_date"):
        profile_text="\n사용자 프로필(개인화에 도움이 될 때만 사용):\n- 이름: "+str(p.get("name",""))+"\n- 자기소개: "+str(p.get("bio",""))+"\n- 생일: "+str(p.get("birth_date",""))
    att=attachments or []
    attachment_text="\n첨부파일:\n"+"\n".join(f"- {a['name']} ({a['type'] or 'unknown'}, {a['size']} bytes)"+("\n  추출된 텍스트:\n"+a['text'][:16000] if a.get('text') else "\n  이 파일은 텍스트를 추출하지 않았습니다.") for a in att) if att else ""
    chart_rule="\n차트 규칙: 사용자가 숫자 데이터를 차트/그래프로 보여달라고 하거나 적절한 시각화를 명시적으로 원하면, 짧은 설명 뒤에 반드시 ```mirae-chart 형태의 JSON 블록을 하나 출력하세요. 형식은 {\"type\":\"bar|line|doughnut\",\"title\":\"제목\",\"labels\":[\"A\",\"B\"],\"datasets\":[{\"label\":\"값\",\"data\":[10,20]}]} 입니다. 데이터가 여러 계열이면 datasets를 여러 개 사용하세요. 숫자가 아닌 내용은 차트로 억지로 만들지 마세요."
    self_info="Mirae AI service facts: public web app domain https://mirae.koharu.live; public API base https://api.koharu.live/v1; backend is FastAPI; PostgreSQL is Neon; generation runs through the Mirae AI model server using the configured local Qwen model. Features: account login, email verification, profile, personalization, themes, selective web search, streaming responses, Markdown/code rendering, account-scoped API keys, HTTP skills, user memory, file attachments, and automatic charts/graphs. API keys are account-scoped and stored hashed. Email verification codes expire after 5 minutes. Never claim the app has capabilities that are not listed here. This self-information is product configuration, not a substitute for live web search."
    return f"""You are Mirae AI, a general-purpose generative AI assistant. Current date: 2026-09-30. {rule}
Do not reveal private chain-of-thought or hidden reasoning. The UI may show only short, high-level progress labels. If the user writes in Korean or Japanese, answer in that language even when the message contains English product names, programming terms, or code. Never switch to English merely because words like discord.py, Python, API, OpenAI, or JavaScript appear. When providing code, keep code in fenced Markdown blocks and keep the surrounding explanation in the user's language. Do not escape Markdown punctuation with backslashes unless the user explicitly asks for literal Markdown source. When the user asks for a chart or graph, use the special fenced block ```mirae-chart with JSON fields type (bar, line, or doughnut), title, labels, and datasets; do not put the chart data into a normal code block. Attachments may contain extracted text; use that text when relevant.
Web search is performed selectively. Do not search for casual conversation, greetings, or ordinary questions that do not require current information. Search when the user explicitly asks to search/find/check sources or when the question clearly depends on current or time-sensitive information. When results are relevant, use only facts directly supported by the provided title, publication date, URL, and snippet. Never fill missing details from memory and never invent a source, quote, statistic, model, date, product release, policy, or link. Treat claims inside a news article as claims by that article unless a primary source is also supplied. Prefer a compact bullet summary over a large table unless the user explicitly asks for a table. Do not present a table unless the supplied source material supports every cell. If the preview is insufficient, say so. Use Markdown for structure when helpful: headings, bullets, numbered lists, emphasis, links, and fenced code blocks with a language tag. When giving code, place it in a fenced code block and do not escape it into a single long line.
Use web results only when they are supplied and do not invent citations. Personality setting: {styles.get(req.personality,"균형 잡힌 말투")}. Follow that style consistently. The user's personal instructions below are active instructions for every answer and must be followed unless they conflict with safety or higher-priority instructions.
Product self-knowledge: {self_info}{profile_text}\nUser instructions: {req.instructions[:4000] or 'none'}.{mem}{attachment_text}{chart_rule}{src}{sk}"""

def ensure_conversation(uid,cid,title="새 대화"):
    if not cid:cid=secrets.token_hex(16)
    with db() as c:
        c.execute("INSERT INTO mirae_conversations(id,user_id,title) VALUES (%s,%s,%s) ON CONFLICT(id) DO NOTHING",[cid,uid,title])
        c.commit()
    return cid

def make_conversation_title(message):
    text=re.sub(r"\s+"," ",str(message or "")).strip()
    text=re.sub(r"^#+\s*","",text)
    text=re.sub(r"^\s*[>*`-]+\s*","",text)
    return text[:60].rstrip() or "새 대화"

async def generate_title(message):
    text=re.sub(r"\s+"," ",str(message or "")).strip()
    if not text:return "새 대화"
    try:
        msgs=[
            {"role":"system","content":"대화 제목을 아주 짧게 요약하세요. 사용자의 요청 핵심만 3~8단어로 표현하고 질문형 문장이나 설명을 만들지 마세요. 이모지와 특수 기호를 쓰지 말고 제목만 출력하세요."},
            {"role":"user","content":text[:2000]}
        ]
        raw=await generate_once(msgs,0.1,24)
        title=re.sub(r"[^0-9A-Za-z가-힣ぁ-ゖァ-ヺ ]"," ",str(raw or ""))
        title=re.sub(r"\s+"," ",title).strip()
        if title:return title[:32].rstrip()
    except Exception:
        pass
    return make_conversation_title(text)[:32]

def save_chat(uid,msg,reply,mode,sources,cid="",attachments=None):
    cid=ensure_conversation(uid,cid)
    att=json.dumps(attachments or [],ensure_ascii=False)
    with db() as c:
        c.execute("INSERT INTO mirae_chat_history(user_id,role,content,mode,model,sources,conversation_id,attachments) VALUES (%s,'user',%s,%s,%s,%s,%s,%s),(%s,'assistant',%s,%s,%s,%s,%s,%s)",[uid,msg,mode,MODEL_NAME,json.dumps(sources,ensure_ascii=False),cid,att,uid,reply,mode,MODEL_NAME,json.dumps(sources,ensure_ascii=False),cid,"[]"])
        c.execute("UPDATE mirae_conversations SET updated_at=now() WHERE id=%s AND user_id=%s",[cid,uid])
        c.commit()
    return cid

def current_title_missing(uid,cid):
    if not cid:return False
    with db() as c:
        row=c.execute("SELECT title FROM mirae_conversations WHERE id=%s AND user_id=%s",[cid,uid]).fetchone()
    return bool(row and row["title"]=="새 대화")

async def finalize_conversation(uid,cid,first_message):
    if not cid:return
    with db() as c:
        row=c.execute("SELECT title FROM mirae_conversations WHERE id=%s AND user_id=%s",[cid,uid]).fetchone()
    if row and row["title"]=="새 대화":
        title=await generate_title(first_message)
        with db() as c:
            c.execute("UPDATE mirae_conversations SET title=%s,updated_at=now() WHERE id=%s AND user_id=%s",[title,cid,uid]);c.commit()

def set_session(resp,uid):
    tok=secrets.token_urlsafe(48); exp=datetime.now(timezone.utc)+timedelta(days=SESSION_DAYS)
    with db() as c:c.execute("INSERT INTO mirae_sessions(user_id,token_hash,expires_at) VALUES (%s,%s,%s)",[uid,digest(tok),exp]);c.commit()
    resp.set_cookie("mirae_session",tok,max_age=SESSION_DAYS*86400,httponly=True,secure=True,samesite="none",path="/")

def init_db():
    if not DATABASE_URL:return
    with db() as c:
        c.execute("ALTER TABLE mirae_users ADD COLUMN IF NOT EXISTS bio TEXT NOT NULL DEFAULT ''")
        c.execute("ALTER TABLE mirae_users ADD COLUMN IF NOT EXISTS birth_date DATE")
        c.execute("ALTER TABLE mirae_users ADD COLUMN IF NOT EXISTS avatar_url TEXT NOT NULL DEFAULT ''")
        c.execute("ALTER TABLE mirae_users ADD COLUMN IF NOT EXISTS email_verified BOOLEAN NOT NULL DEFAULT true")
        c.execute("ALTER TABLE mirae_chat_history ADD COLUMN IF NOT EXISTS conversation_id TEXT NOT NULL DEFAULT ''")
        c.execute("ALTER TABLE mirae_chat_history ADD COLUMN IF NOT EXISTS attachments JSONB NOT NULL DEFAULT '[]'::jsonb")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_email_verifications(
            id BIGSERIAL PRIMARY KEY,email TEXT NOT NULL,name TEXT NOT NULL,password_hash TEXT NOT NULL,
            code_hash TEXT NOT NULL,expires_at TIMESTAMPTZ NOT NULL,attempts INTEGER NOT NULL DEFAULT 0,created_at TIMESTAMPTZ NOT NULL DEFAULT now())""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_email_verifications_email_idx ON mirae_email_verifications(email,created_at DESC)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_skills(
            id BIGSERIAL PRIMARY KEY,user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            name TEXT NOT NULL,description TEXT NOT NULL DEFAULT '',url TEXT NOT NULL,method TEXT NOT NULL DEFAULT 'POST',
            headers TEXT NOT NULL DEFAULT '',body TEXT NOT NULL DEFAULT '',active BOOLEAN NOT NULL DEFAULT true,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),UNIQUE(user_id,name))""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_skills_user_idx ON mirae_skills(user_id)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_plugins(
            id BIGSERIAL PRIMARY KEY,user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            name TEXT NOT NULL,description TEXT NOT NULL DEFAULT '',url TEXT NOT NULL,
            method TEXT NOT NULL DEFAULT 'POST',permissions JSONB NOT NULL DEFAULT '[]'::jsonb,
            active BOOLEAN NOT NULL DEFAULT true,created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),UNIQUE(user_id,name)
        )""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_plugins_user_idx ON mirae_plugins(user_id)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_conversations(
            id TEXT PRIMARY KEY,
            user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            title TEXT NOT NULL DEFAULT '새 대화',
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
        )""")
        c.execute("ALTER TABLE mirae_conversations ADD COLUMN IF NOT EXISTS favorite BOOLEAN NOT NULL DEFAULT false")
        c.execute("ALTER TABLE mirae_conversations ADD COLUMN IF NOT EXISTS folder_id BIGINT")
        c.execute("ALTER TABLE mirae_conversations ADD COLUMN IF NOT EXISTS share_code TEXT")
        c.execute("CREATE UNIQUE INDEX IF NOT EXISTS mirae_conversations_share_idx ON mirae_conversations(share_code) WHERE share_code IS NOT NULL")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_conversation_folders(
            id BIGSERIAL PRIMARY KEY,
            user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            name TEXT NOT NULL,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            UNIQUE(user_id,name)
        )""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_conversation_folders_user_idx ON mirae_conversation_folders(user_id)")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_conversations_user_idx ON mirae_conversations(user_id,updated_at DESC)")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_chat_history_user_created_idx ON mirae_chat_history(user_id,created_at DESC)")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_chat_history_conversation_idx ON mirae_chat_history(user_id,conversation_id,created_at ASC,id ASC)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_memories(
            id BIGSERIAL PRIMARY KEY,
            user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            content TEXT NOT NULL,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
        )""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_memories_user_idx ON mirae_memories(user_id,updated_at DESC)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_api_usage(
            id BIGSERIAL PRIMARY KEY,key_id BIGINT REFERENCES mirae_api_keys(id) ON DELETE SET NULL,
            user_id BIGINT REFERENCES mirae_users(id) ON DELETE CASCADE,path TEXT NOT NULL,
            status INTEGER NOT NULL DEFAULT 200,latency_ms INTEGER NOT NULL DEFAULT 0,
            prompt_tokens INTEGER NOT NULL DEFAULT 0,completion_tokens INTEGER NOT NULL DEFAULT 0,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now())""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_api_usage_user_idx ON mirae_api_usage(user_id,created_at DESC)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_webhooks(
            id BIGSERIAL PRIMARY KEY,user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            name TEXT NOT NULL,url TEXT NOT NULL,events JSONB NOT NULL DEFAULT '["api.request"]'::jsonb,
            active BOOLEAN NOT NULL DEFAULT true,created_at TIMESTAMPTZ NOT NULL DEFAULT now())""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_webhooks_user_idx ON mirae_webhooks(user_id)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_access_logs(
            id BIGSERIAL PRIMARY KEY,user_id BIGINT REFERENCES mirae_users(id) ON DELETE SET NULL,
            email TEXT NOT NULL DEFAULT '',nickname TEXT NOT NULL DEFAULT '',ip TEXT NOT NULL DEFAULT '',
            country TEXT NOT NULL DEFAULT '',region TEXT NOT NULL DEFAULT '',city TEXT NOT NULL DEFAULT '',
            path TEXT NOT NULL,method TEXT NOT NULL,status INTEGER NOT NULL,
            risk_category TEXT NOT NULL DEFAULT 'normal',risk_detail TEXT NOT NULL DEFAULT '',
            user_agent TEXT NOT NULL DEFAULT '',latency_ms INTEGER NOT NULL DEFAULT 0,created_at TIMESTAMPTZ NOT NULL DEFAULT now())""")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_access_logs_created_idx ON mirae_access_logs(created_at DESC)")
        c.execute("CREATE INDEX IF NOT EXISTS mirae_access_logs_risk_idx ON mirae_access_logs(risk_category,created_at DESC)")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_admin_settings(key TEXT PRIMARY KEY,value JSONB NOT NULL DEFAULT '{}'::jsonb,updated_at TIMESTAMPTZ NOT NULL DEFAULT now())""")
        c.execute("""CREATE TABLE IF NOT EXISTS mirae_feedback(
            id BIGSERIAL PRIMARY KEY,
            user_id BIGINT NOT NULL REFERENCES mirae_users(id) ON DELETE CASCADE,
            conversation_id TEXT,
            message_hash TEXT NOT NULL,
            feedback TEXT NOT NULL CHECK(feedback IN ('like','dislike')),
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            UNIQUE(user_id,message_hash)
        )""")
        c.commit()


def extract_file_text(name,data):
    ext=os.path.splitext(name.lower())[1]
    try:
        if ext==".pdf":
            from pypdf import PdfReader
            r=PdfReader(BytesIO(data));return "\n".join((p.extract_text() or "") for p in r.pages)
        if ext==".docx":
            from docx import Document
            d=Document(BytesIO(data));return "\n".join(x.text for x in d.paragraphs)
        if ext==".xlsx":
            from openpyxl import load_workbook
            wb=load_workbook(BytesIO(data),read_only=True,data_only=True);parts=[]
            for ws in wb.worksheets:
                parts.append("[Sheet: "+ws.title+"]")
                for row in ws.iter_rows(values_only=True):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(parts)
        if ext==".pptx":
            from pptx import Presentation
            prs=Presentation(BytesIO(data));parts=[]
            for i,slide in enumerate(prs.slides,1):
                parts.append("[Slide "+str(i)+"]")
                for shape in slide.shapes:
                    if hasattr(shape,"text"):parts.append(shape.text)
            return "\n".join(parts)
        if ext==".zip":
            parts=[]
            with zipfile.ZipFile(BytesIO(data)) as z:
                for n in z.namelist():
                    if n.endswith("/") or len(parts)>=200:continue
                    parts.append("[FILE] "+n)
                    if z.getinfo(n).file_size<=500_000 and re.search(r"\.(txt|md|json|csv|py|js|ts|tsx|jsx|html|css|sql|yaml|yml|xml|toml|ini|log)$",n,re.I):
                        try:parts.append(z.read(n).decode("utf-8","ignore")[:20000])
                        except:pass
            return "\n".join(parts)
        if ext in {".txt",".md",".csv",".json",".py",".js",".ts",".tsx",".jsx",".html",".css",".sql",".yaml",".yml",".xml",".toml",".ini",".log"}:
            return data.decode("utf-8","ignore")
    except Exception as e:
        return "[파일 텍스트 추출 실패: "+str(e)[:200]+"]"
    return ""

@app.post("/files/extract")
async def extract_files(files:list[UploadFile]=File(...)):
    if len(files)>20:raise HTTPException(400,"한 번에 최대 20개 파일까지 처리할 수 있습니다.")
    out=[]
    for f in files:
        data=await f.read()
        if len(data)>100*1024*1024:raise HTTPException(413,f"{f.filename} 파일이 100MB를 초과합니다.")
        text=await asyncio.to_thread(extract_file_text,f.filename or "file",data)
        # 대용량 파일은 청크 단위로 잘라 모델 컨텍스트를 보호합니다.
        chunks=[text[i:i+16000] for i in range(0,len(text),16000)] or [""]
        out.append({"name":f.filename or "file","type":f.content_type or "application/octet-stream","size":len(data),"text":text[:200000],"chunks":chunks[:100],"chunk_count":len(chunks),"truncated":len(text)>200000})
    return {"files":out}

async def translate_image_prompt(prompt:str)->str:
    text=str(prompt or "").strip()
    if not text or not re.search(r"[가-힣]",text):return text
    try:
        from googletrans import Translator
        async with Translator(service_urls=["translate.googleapis.com"]) as translator:
            translated=await translator.translate(text,dest="en")
        return str(translated.text or text).strip() or text
    except Exception:
        return text

async def generate_image_payload(prompt:str):
    prompt=str(prompt or "").strip()
    if not prompt: raise HTTPException(400,"이미지 프롬프트가 필요합니다.")
    if len(prompt)>4000: raise HTTPException(400,"이미지 프롬프트는 4000자까지 입력할 수 있습니다.")
    if not CLOUDFLARE_ACCOUNT_ID or not CLOUDFLARE_API_TOKEN: raise HTTPException(503,"이미지 생성 서버가 설정되지 않았습니다.")
    translated_prompt=await translate_image_prompt(prompt)
    payload={"prompt":translated_prompt,"num_steps":20,"guidance":7.5,"width":1024,"height":1024}
    url="https://api.cloudflare.com/client/v4/accounts/"+CLOUDFLARE_ACCOUNT_ID+"/ai/run/"+CLOUDFLARE_IMAGE_MODEL
    try:
        async with httpx.AsyncClient(timeout=180,trust_env=False) as x:
            r=await x.post(url,headers={"Authorization":"Bearer "+CLOUDFLARE_API_TOKEN,"Content-Type":"application/json"},json=payload)
            if r.status_code>=400: raise HTTPException(502,"Cloudflare 이미지 생성 서비스 요청이 거부되었습니다.")
            image_bytes=r.content if "image/" in (r.headers.get("content-type") or "").lower() else b""
            if not image_bytes:
                try:
                    body=r.json(); result=body.get("result") if isinstance(body,dict) else None
                    if isinstance(result,str): image_bytes=base64.b64decode(result)
                    elif isinstance(result,dict):
                        encoded=result.get("image") or result.get("image_b64") or result.get("b64_json")
                        if encoded:image_bytes=base64.b64decode(encoded)
                except Exception: pass
            if not image_bytes: raise HTTPException(502,"Cloudflare 이미지 생성 결과가 없습니다.")
            return {"ok":True,"url":"data:image/png;base64,"+base64.b64encode(image_bytes).decode(),"model":CLOUDFLARE_IMAGE_MODEL,"prompt":prompt,"translated_prompt":translated_prompt}
    except HTTPException: raise
    except Exception as e: raise HTTPException(502,"Cloudflare 이미지 생성 요청에 실패했습니다: "+type(e).__name__)
def classify_image_intent(message:str):
    text=str(message or "").strip()
    if not text:return {"generate":False,"prompt":""}
    terms=("이미지","그림","일러스트","사진","포스터","아이콘","로고")
    create=("만들","생성","그려","제작")
    if not any(a in text for a in terms) or not any(b in text for b in create):
        return {"generate":False,"prompt":""}
    return {"generate":True,"prompt":text[:4000]}

@app.post("/images/generate")
async def generate_image(data:dict[str,Any],request:Request):
    if not session_user(request): raise HTTPException(401,"로그인이 필요합니다.")
    return await generate_image_payload(str(data.get("prompt","")))
@app.post("/vision")
async def understand_image(file:UploadFile=File(...),request:Request=None):
    if not session_user(request):raise HTTPException(401,"로그인이 필요합니다.")
    if not (file.content_type or "").startswith("image/"):raise HTTPException(400,"이미지 파일만 사용할 수 있습니다.")
    data=await file.read()
    if len(data)>12*1024*1024:raise HTTPException(413,"이미지는 12MB 이하만 분석할 수 있습니다.")
    if not POLLINATIONS_API_KEY:raise HTTPException(503,"비전 모델이 설정되지 않았습니다.")
    data_url="data:"+(file.content_type or "image/png")+";base64,"+base64.b64encode(data).decode()
    messages=[{"role":"user","content":[{"type":"text","text":"이 이미지를 한국어로 정확하게 분석해줘. 보이는 사실과 텍스트를 구분하고, 보이지 않는 내용은 추측하지 마."},{"type":"image_url","image_url":{"url":data_url}}]}]
    try:
        async with httpx.AsyncClient(timeout=90,trust_env=False) as x:
            r=await x.post("https://gen.pollinations.ai/v1/chat/completions",headers={"Authorization":"Bearer "+POLLINATIONS_API_KEY,"Content-Type":"application/json"},json={"model":POLLINATIONS_VISION_MODEL,"messages":messages,"max_tokens":1200})
            if r.status_code>=400:raise HTTPException(502,"비전 모델이 이미지를 처리하지 못했습니다.")
            body=r.json();content=((body.get("choices") or [{}])[0].get("message") or {}).get("content","")
            if not content:raise HTTPException(502,"비전 모델이 분석 결과를 반환하지 않았습니다.")
            return {"ok":True,"text":str(content)[:20000],"model":POLLINATIONS_VISION_MODEL}
    except HTTPException:raise
    except Exception as e:raise HTTPException(502,"이미지 분석 요청에 실패했습니다: "+type(e).__name__)

@app.on_event("startup")
async def startup():
    try:init_db()
    except Exception:pass

def send_code(to,code):
    if not RESEND_API_KEY:raise RuntimeError("RESEND_API_KEY is not configured.")
    payload={
        "from":RESEND_FROM,
        "to":[to],
        "subject":"[Mirae AI] 이메일 인증 코드",
        "text":f"Mirae AI 인증 코드: {code}\n\n이 코드는 5분 후 만료됩니다.",
        "html":f"<div style='font-family:Arial,sans-serif;padding:30px'><h2>Mirae AI 이메일 인증</h2><p>인증 코드를 입력하세요.</p><div style='font-size:32px;font-weight:700;letter-spacing:8px'>{code}</div><p>이 코드는 <b>5분 후 만료</b>됩니다.</p></div>",
        "tags":[{"name":"category","value":"confirm_email"}]
    }
    with httpx.Client(timeout=20) as x:
        r=x.post(RESEND_URL,headers={"Authorization":f"Bearer {RESEND_API_KEY}","Content-Type":"application/json"},json=payload)
    if r.status_code>=400:
        detail=r.text.replace(RESEND_API_KEY,"[REDACTED]")[:500]
        raise RuntimeError(f"Resend API {r.status_code}: {detail}")

async def request_signup(data):
    email=data.email.strip().lower()
    if "@" not in email:raise HTTPException(400,"올바른 이메일을 입력해주세요.")
    with db() as c:
        if c.execute("SELECT id FROM mirae_users WHERE email=%s",[email]).fetchone():raise HTTPException(409,"이미 가입된 이메일입니다.")
        r=c.execute("SELECT created_at FROM mirae_email_verifications WHERE email=%s ORDER BY created_at DESC LIMIT 1",[email]).fetchone()
        if r and r["created_at"]>datetime.now(timezone.utc)-timedelta(seconds=60):raise HTTPException(429,"인증 코드는 60초마다 다시 요청할 수 있습니다.")
        code=f"{secrets.randbelow(1000000):06d}"
        c.execute("DELETE FROM mirae_email_verifications WHERE email=%s",[email])
        c.execute("INSERT INTO mirae_email_verifications(email,name,password_hash,code_hash,expires_at) VALUES (%s,%s,%s,%s,%s)",[email,data.name.strip(),phash(data.password),digest(code),datetime.now(timezone.utc)+timedelta(minutes=5)]);c.commit()
    try:await asyncio.to_thread(send_code,email,code)
    except Exception as e:
        with db() as c:c.execute("DELETE FROM mirae_email_verifications WHERE email=%s",[email]);c.commit()
        detail=str(e).replace(RESEND_API_KEY,"[REDACTED]")[:500]
        print(f"EMAIL_SEND_ERROR provider=resend_api from={RESEND_FROM}: {detail}",flush=True)
        raise HTTPException(503,f"인증 메일을 보내지 못했습니다. Resend 오류: {type(e).__name__}: {detail}")
    return {"verification_required":True,"expires_in":300}

@app.get("/docs",include_in_schema=False)
async def custom_docs():
    return RedirectResponse("https://mirae.koharu.live/api-docs.html")

@app.get("/health")
async def health():
    return {"ok":True,"model":MODEL_NAME,"web_search":True,"database":bool(DATABASE_URL),"version":APP_VERSION,"email_verification":bool(RESEND_API_KEY),"email_provider":"resend_api"}

@app.post("/auth/signup")
async def signup(data:Signup):return await request_signup(data)
@app.post("/auth/signup/request")
async def signup_request(data:Signup):return await request_signup(data)

@app.post("/auth/signup/verify")
async def signup_verify(data:Verify,response:Response):
    email=data.email.strip().lower()
    with db() as c:
        row=c.execute("SELECT * FROM mirae_email_verifications WHERE email=%s ORDER BY created_at DESC LIMIT 1",[email]).fetchone()
        if not row:raise HTTPException(404,"인증 요청을 찾을 수 없습니다. 다시 요청해주세요.")
        if row["expires_at"]<=datetime.now(timezone.utc):
            c.execute("DELETE FROM mirae_email_verifications WHERE id=%s",[row["id"]]);c.commit();raise HTTPException(410,"인증 코드가 만료되었습니다.")
        if row["attempts"]>=5:raise HTTPException(429,"인증 시도 횟수를 초과했습니다.")
        if not hmac.compare_digest(digest(data.code),row["code_hash"]):
            c.execute("UPDATE mirae_email_verifications SET attempts=attempts+1 WHERE id=%s",[row["id"]]);c.commit();raise HTTPException(400,"인증 코드가 올바르지 않습니다.")
        if c.execute("SELECT id FROM mirae_users WHERE email=%s",[email]).fetchone():raise HTTPException(409,"이미 가입된 이메일입니다.")
        u=c.execute("INSERT INTO mirae_users(email,password_hash,name,email_verified) VALUES (%s,%s,%s,true) RETURNING id,email,name",[email,row["password_hash"],row["name"]]).fetchone()
        c.execute("INSERT INTO mirae_user_settings(user_id) VALUES (%s)",[u["id"]]);c.execute("DELETE FROM mirae_email_verifications WHERE email=%s",[email]);c.commit()
    set_session(response,u["id"]);return {"user":u,"verified":True}

@app.post("/auth/login")
async def login(data:Login,response:Response):
    with db() as c:u=c.execute("SELECT * FROM mirae_users WHERE email=%s",[data.email.strip().lower()]).fetchone()
    if not u or not pok(data.password,u["password_hash"]):raise HTTPException(401,"이메일 또는 비밀번호가 올바르지 않습니다.")
    if not u.get("email_verified",True):raise HTTPException(403,"이메일 인증이 완료되지 않은 계정입니다.")
    set_session(response,u["id"])
    with db() as c:c.execute("UPDATE mirae_users SET last_signed_in=now(),updated_at=now() WHERE id=%s",[u["id"]]);c.commit()
    return {"user":{"id":u["id"],"email":u["email"],"name":u["name"]}}

@app.post("/auth/logout")
async def logout(request:Request,response:Response):
    t=request.cookies.get("mirae_session")
    if t:
        with db() as c:c.execute("DELETE FROM mirae_sessions WHERE token_hash=%s",[digest(t)]);c.commit()
    response.delete_cookie("mirae_session",path="/");return {"ok":True}

@app.get("/auth/me")
async def me(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    return {"user":{"id":u["id"],"email":u["email"],"name":u["name"]}}

@app.get("/profile")
async def get_profile(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    return {"name":u["name"],"email":u["email"],"bio":u.get("bio",""),"birth_date":u["birth_date"].isoformat() if u.get("birth_date") else None,"avatar_url":u.get("avatar_url","")}

@app.put("/profile")
async def put_profile(data:Profile,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    if data.avatar_url and not re.match(r"^https?://",data.avatar_url,re.I):raise HTTPException(400,"프로필 이미지 URL은 http:// 또는 https://여야 합니다.")
    with db() as c:
        r=c.execute("UPDATE mirae_users SET name=%s,bio=%s,birth_date=%s,avatar_url=%s,updated_at=now() WHERE id=%s RETURNING id,email,name,bio,birth_date,avatar_url",[data.name.strip(),data.bio.strip(),data.birth_date,data.avatar_url.strip(),u["id"]]).fetchone();c.commit()
    return {"name":r["name"],"email":r["email"],"bio":r["bio"],"birth_date":r["birth_date"].isoformat() if r["birth_date"] else None,"avatar_url":r["avatar_url"]}

@app.get("/settings")
async def get_settings(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:r=c.execute("SELECT theme,personality,instructions,web_search,temperature FROM mirae_user_settings WHERE user_id=%s",[u["id"]]).fetchone()
    return r or {"theme":"light","personality":"balanced","instructions":"","web_search":True,"temperature":.7}

@app.put("/settings")
async def put_settings(data:Settings,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    personality=data.personality if data.personality in {"balanced","friendly","professional","concise","creative"} else "balanced"
    instructions=data.instructions.strip()[:4000]
    with db() as c:
        c.execute("""INSERT INTO mirae_user_settings(user_id,theme,personality,instructions,web_search,temperature)
                     VALUES (%s,%s,%s,%s,%s,%s)
                     ON CONFLICT(user_id) DO UPDATE SET theme=EXCLUDED.theme,personality=EXCLUDED.personality,
                     instructions=EXCLUDED.instructions,web_search=EXCLUDED.web_search,temperature=EXCLUDED.temperature,updated_at=now()""",
                  [u["id"],data.theme,personality,instructions,bool(data.web_search),data.temperature]);c.commit()
    result={"theme":data.theme,"personality":personality,"instructions":instructions,"web_search":bool(data.web_search),"temperature":data.temperature}
    return result

@app.get("/history")
async def history(request:Request,limit:int=200):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    limit=max(1,min(limit,500))
    with db() as c:
        r=c.execute("""SELECT h.id,h.role,h.content,h.mode,h.model,h.sources,h.conversation_id,h.attachments,h.created_at,
                              COALESCE(cv.title,'새 대화') AS conversation_title
                       FROM mirae_chat_history h
                       LEFT JOIN mirae_conversations cv ON cv.id=h.conversation_id
                       WHERE h.user_id=%s ORDER BY h.created_at DESC,h.id DESC LIMIT %s""",[u["id"],limit]).fetchall()
    return list(reversed(r))

@app.get("/conversations")
async def conversations(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        rows=c.execute("""SELECT c.id,c.title,c.created_at,c.updated_at,c.favorite,c.folder_id,c.share_code,
                                 first_msg.content AS first_user_message
                          FROM mirae_conversations c
                          LEFT JOIN LATERAL (
                              SELECT h.content
                              FROM mirae_chat_history h
                              WHERE h.user_id=c.user_id AND h.conversation_id=c.id AND h.role='user'
                              ORDER BY h.created_at ASC,h.id ASC LIMIT 1
                          ) first_msg ON true
                          WHERE c.user_id=%s
                          ORDER BY c.favorite DESC,c.updated_at DESC LIMIT 500""",[u["id"]]).fetchall()
    out=[]
    for row in rows:
        item=dict(row)
        title=str(item.get("title") or "")
        if title=="새 대화" or ("@" in title and "." in title):
            item["title"]=make_conversation_title(item.get("first_user_message",""))
        item.pop("first_user_message",None)
        out.append(item)
    return out

@app.get("/conversations/{conversation_id}/messages")
async def conversation_messages(conversation_id:str,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        exists=c.execute("SELECT id FROM mirae_conversations WHERE id=%s AND user_id=%s",[conversation_id,u["id"]]).fetchone()
        if not exists:raise HTTPException(404,"대화를 찾을 수 없습니다.")
        rows=c.execute("""SELECT id,role,content,sources,attachments,created_at
                          FROM mirae_chat_history
                          WHERE user_id=%s AND conversation_id=%s
                          ORDER BY created_at ASC,id ASC""",[u["id"],conversation_id]).fetchall()
    return rows

@app.get("/conversations/search")
async def search_conversations(q:str="",request:Request=None):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    query=re.sub(r"\s+"," ",str(q or "")).strip()[:120]
    if not query:return await conversations(request)
    like="%"+query.replace("%","\\%").replace("_","\\_")+"%"
    with db() as c:
        rows=c.execute("""SELECT id,title,created_at,updated_at,favorite,folder_id,share_code
                          FROM mirae_conversations
                          WHERE user_id=%s AND (title ILIKE %s ESCAPE '\' OR id IN
                            (SELECT conversation_id FROM mirae_chat_history WHERE user_id=%s AND content ILIKE %s ESCAPE '\'))
                          ORDER BY favorite DESC,updated_at DESC LIMIT 100""",[u["id"],like,u["id"],like]).fetchall()
    return rows

@app.get("/conversation-folders")
async def list_conversation_folders(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:return c.execute("SELECT id,name,created_at FROM mirae_conversation_folders WHERE user_id=%s ORDER BY name ASC",[u["id"]]).fetchall()

@app.post("/conversation-folders")
async def create_conversation_folder(data:ConversationFolder,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    name=re.sub(r"\s+"," ",data.name.strip())[:60]
    with db() as c:
        try:r=c.execute("INSERT INTO mirae_conversation_folders(user_id,name) VALUES (%s,%s) RETURNING id,name",[u["id"],name]).fetchone()
        except psycopg.errors.UniqueViolation:raise HTTPException(409,"같은 이름의 폴더가 이미 있습니다.")
        c.commit()
    return r

@app.delete("/conversation-folders/{folder_id}")
async def delete_conversation_folder(folder_id:int,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        c.execute("UPDATE mirae_conversations SET folder_id=NULL WHERE user_id=%s AND folder_id=%s",[u["id"],folder_id])
        r=c.execute("DELETE FROM mirae_conversation_folders WHERE id=%s AND user_id=%s RETURNING id",[folder_id,u["id"]]).fetchone();c.commit()
    return {"ok":True,"deleted":bool(r)}

@app.put("/conversations/{conversation_id}/favorite")
async def favorite_conversation(conversation_id:str,data:ConversationFavorite,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        r=c.execute("UPDATE mirae_conversations SET favorite=%s,updated_at=now() WHERE id=%s AND user_id=%s RETURNING id,favorite",[data.favorite,conversation_id,u["id"]]).fetchone()
        if not r:raise HTTPException(404,"대화를 찾을 수 없습니다.")
        c.commit()
    return r

@app.put("/conversations/{conversation_id}/folder")
async def assign_conversation_folder(conversation_id:str,data:ConversationFolderAssign,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        if data.folder_id is not None:
            ok=c.execute("SELECT id FROM mirae_conversation_folders WHERE id=%s AND user_id=%s",[data.folder_id,u["id"]]).fetchone()
            if not ok:raise HTTPException(404,"폴더를 찾을 수 없습니다.")
        r=c.execute("UPDATE mirae_conversations SET folder_id=%s,updated_at=now() WHERE id=%s AND user_id=%s RETURNING id,folder_id",[data.folder_id,conversation_id,u["id"]]).fetchone()
        if not r:raise HTTPException(404,"대화를 찾을 수 없습니다.")
        c.commit()
    return r

@app.post("/conversations/{conversation_id}/share")
async def share_conversation(conversation_id:str,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        row=c.execute("SELECT id,share_code FROM mirae_conversations WHERE id=%s AND user_id=%s",[conversation_id,u["id"]]).fetchone()
        if not row:raise HTTPException(404,"대화를 찾을 수 없습니다.")
        code=row["share_code"]
        if not code:
            alphabet="ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"
            while True:
                code="".join(secrets.choice(alphabet) for _ in range(8))
                if not c.execute("SELECT 1 FROM mirae_conversations WHERE share_code=%s",[code]).fetchone():break
            c.execute("UPDATE mirae_conversations SET share_code=%s WHERE id=%s AND user_id=%s",[code,conversation_id,u["id"]]);c.commit()
    return {"ok":True,"code":code,"url":"https://mirae.koharu.live/share/"+code}

@app.get("/shared/{share_code}")
async def get_shared_conversation(share_code:str):
    if not re.fullmatch(r"[A-Za-z]{8}",share_code):raise HTTPException(404,"공유 대화를 찾을 수 없습니다.")
    with db() as c:
        conv=c.execute("SELECT id,title,created_at,updated_at FROM mirae_conversations WHERE share_code=%s",[share_code]).fetchone()
        if not conv:raise HTTPException(404,"공유 대화를 찾을 수 없습니다.")
        messages=c.execute("SELECT role,content,created_at,attachments FROM mirae_chat_history WHERE conversation_id=%s ORDER BY created_at ASC,id ASC",[conv["id"]]).fetchall()
    return {"conversation":conv,"messages":messages}

@app.put("/conversations/{conversation_id}")
async def rename_conversation(conversation_id:str,data:ConversationRename,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    title=re.sub(r"\s+"," ",data.title.strip())[:120]
    with db() as c:
        r=c.execute("UPDATE mirae_conversations SET title=%s,updated_at=now() WHERE id=%s AND user_id=%s RETURNING id,title",[title,conversation_id,u["id"]]).fetchone()
        if not r:raise HTTPException(404,"대화를 찾을 수 없습니다.")
        c.commit()
    return r

@app.delete("/conversations/{conversation_id}")
async def delete_conversation(conversation_id:str,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        c.execute("DELETE FROM mirae_chat_history WHERE user_id=%s AND conversation_id=%s",[u["id"],conversation_id])
        r=c.execute("DELETE FROM mirae_conversations WHERE user_id=%s AND id=%s RETURNING id",[u["id"],conversation_id]).fetchone()
        c.commit()
    return {"ok":True,"deleted":bool(r)}

def auto_register_memory(uid:int,message:str):
    text=re.sub(r"\s+"," ",str(message or "")).strip()
    if not text:return
    candidates=[]
    if re.search(r"(기억해|기억하|잊지 말|앞으로|내 이름은|나는|제가|내가)",text,re.I):
        clean=re.sub(r"^(?:기억해줘|기억해|잊지 말아줘|앞으로)\s*[:：]?\s*","",text,flags=re.I).strip()
        if clean:
            candidates.append(clean[:1000])
    for content in candidates:
        with db() as c:
            exists=c.execute("SELECT id FROM mirae_memories WHERE user_id=%s AND lower(content)=lower(%s) LIMIT 1",[uid,content]).fetchone()
            if not exists:
                c.execute("INSERT INTO mirae_memories(user_id,content) VALUES (%s,%s)",[uid,content]);c.commit()


@app.get("/admin/overview")
async def admin_overview(request:Request):
    require_admin(request)
    with db() as c:
        users=c.execute("SELECT COUNT(*) AS count FROM mirae_users").fetchone()["count"]
        active=c.execute("SELECT COUNT(*) AS count FROM mirae_sessions WHERE expires_at>now()").fetchone()["count"]
        logs=c.execute("SELECT COUNT(*) AS count FROM mirae_access_logs WHERE created_at>=now()-interval '24 hours'").fetchone()["count"]
        risks=c.execute("SELECT risk_category,COUNT(*) AS count FROM mirae_access_logs WHERE risk_category<>'normal' AND created_at>=now()-interval '7 days' GROUP BY risk_category ORDER BY count DESC").fetchall()
    return {"users":users,"active_sessions":active,"access_24h":logs,"risks":risks}
@app.get("/admin/users")
async def admin_users(request:Request):
    require_admin(request)
    with db() as c:return c.execute("""SELECT u.id,u.email,u.name,u.created_at,u.email_verified,a.last_access,a.ip,a.country,a.region,a.city FROM mirae_users u LEFT JOIN LATERAL (SELECT created_at AS last_access,ip,country,region,city FROM mirae_access_logs l WHERE l.user_id=u.id ORDER BY created_at DESC LIMIT 1) a ON true ORDER BY u.created_at DESC LIMIT 500""").fetchall()
@app.get("/admin/logs")
async def admin_logs(request:Request,category:str="",limit:int=200):
    require_admin(request); limit=max(1,min(limit,500))
    with db() as c:
        if category:return c.execute("SELECT id,email,nickname,ip,country,region,city,path,method,status,risk_category,risk_detail,user_agent,latency_ms,created_at FROM mirae_access_logs WHERE risk_category=%s ORDER BY created_at DESC LIMIT %s",[category,limit]).fetchall()
        return c.execute("SELECT id,email,nickname,ip,country,region,city,path,method,status,risk_category,risk_detail,user_agent,latency_ms,created_at FROM mirae_access_logs ORDER BY created_at DESC LIMIT %s",[limit]).fetchall()
@app.get("/admin/security")
async def admin_security(request:Request):
    require_admin(request)
    now=time.time(); threats=[]
    for ip,items in traffic_snapshot():
        recent=[t for t in items if t>=now-60]
        if len(recent)>=60: threats.append({"ip":ip,"requests":len(recent),"reason":"최근 1분 요청 폭주"})
        elif len(items)>=240: threats.append({"ip":ip,"requests":len(items),"reason":"최근 15분 누적 요청 급증"})
    with db() as c:
        total=c.execute("SELECT COUNT(*) AS n FROM mirae_access_logs WHERE created_at>=now()-interval '15 minutes'").fetchone()["n"]
        errors=c.execute("SELECT COUNT(*) AS n FROM mirae_access_logs WHERE created_at>=now()-interval '15 minutes' AND status>=400").fetchone()["n"]
        limited=c.execute("SELECT COUNT(*) AS n FROM mirae_access_logs WHERE created_at>=now()-interval '15 minutes' AND status=429").fetchone()["n"]
    return {"state":"alert" if threats or (total and errors/total>=0.2) else "normal","suspicious_ips":len(threats),"bursts":len(threats),"error_rate":(errors/total*100 if total else 0),"rate_limited":limited,"threats":threats[:100]}

@app.get("/admin/traffic")
async def admin_traffic(request:Request):
    require_admin(request)
    with db() as c:
        stats=c.execute("""SELECT COUNT(*) AS requests,COALESCE(AVG(latency_ms),0) AS avg_latency,COALESCE(AVG(CASE WHEN status<400 THEN 1.0 ELSE 0.0 END)*100,100) AS success_rate,COUNT(*) FILTER(WHERE status>=500) AS server_errors FROM mirae_access_logs WHERE created_at>=now()-interval '24 hours'""").fetchone()
        endpoints=c.execute("""SELECT path,COUNT(*) AS count FROM mirae_access_logs WHERE created_at>=now()-interval '24 hours' GROUP BY path ORDER BY count DESC LIMIT 20""").fetchall()
    return {**dict(stats),"endpoints":endpoints}

@app.post("/admin/cleanup-logs")
async def admin_cleanup_logs(request:Request):
    require_admin(request)
    with db() as c:
        row=c.execute("SELECT value FROM mirae_admin_settings WHERE key='log_retention_days'").fetchone()
        raw=row["value"] if row else 90
        try: days=int(raw.get("value",raw) if isinstance(raw,dict) else raw)
        except Exception: days=90
        days=max(1,min(days,3650))
        r=c.execute("DELETE FROM mirae_access_logs WHERE created_at<now()-(%s::text||' days')::interval",[days]); c.commit()
    return {"ok":True,"deleted":r.rowcount,"retention_days":days}

@app.get("/admin/settings")
async def admin_settings(request:Request):
    require_admin(request); defaults={"web_search_mode":"broad","image_generation":True,"maintenance":False,"log_retention_days":90,"security_sensitivity":"normal","admin_audit_enabled":True}
    with db() as c:
        for r in c.execute("SELECT key,value FROM mirae_admin_settings ORDER BY key").fetchall(): defaults[r["key"]]=r["value"]
    return defaults
@app.put("/admin/settings")
async def admin_settings_update(data:dict[str,Any],request:Request):
    require_admin(request); allowed={"web_search_mode","image_generation","maintenance","log_retention_days","security_sensitivity","admin_audit_enabled"}
    with db() as c:
        for k,v in data.items():
            if k in allowed:c.execute("""INSERT INTO mirae_admin_settings(key,value,updated_at) VALUES (%s,%s,now()) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value,updated_at=now()""",[k,Jsonb(v)])
        c.commit()
    return await admin_settings(request)
@app.get("/memories")
async def list_memories(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:return c.execute("SELECT id,content,created_at,updated_at FROM mirae_memories WHERE user_id=%s ORDER BY updated_at DESC LIMIT 100",[u["id"]]).fetchall()

@app.post("/memories")
async def create_memory(data:MemoryCreate,request:Request):
    raise HTTPException(403,"메모리는 Mirae AI만 자동으로 등록할 수 있습니다.")

@app.delete("/memories/{memory_id}")
async def delete_memory(memory_id:int,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        r=c.execute("DELETE FROM mirae_memories WHERE id=%s AND user_id=%s RETURNING id",[memory_id,u["id"]]).fetchone();c.commit()
    return {"ok":True,"deleted":bool(r)}

@app.put("/feedback")
async def feedback(data:dict[str,Any],request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    value=str(data.get("feedback",""))
    if value not in {"like","dislike"}:raise HTTPException(400,"feedback must be like or dislike.")
    mh=str(data.get("message_hash","")).strip()
    if not mh:raise HTTPException(400,"message_hash is required.")
    with db() as c:
        c.execute("""INSERT INTO mirae_feedback(user_id,conversation_id,message_hash,feedback)
                     VALUES (%s,%s,%s,%s)
                     ON CONFLICT(user_id,message_hash) DO UPDATE SET feedback=EXCLUDED.feedback,conversation_id=EXCLUDED.conversation_id""",
                  [u["id"],str(data.get("conversation_id","")),mh,value]);c.commit()
    return {"ok":True,"feedback":value}

def templ(v:Any,p:dict):
    if isinstance(v,str):
        def f(m):
            cur=p
            for k in m.group(1).split("."):cur=cur.get(k,"") if isinstance(cur,dict) else ""
            return str(cur)
        return re.sub(r"\{\{\s*([A-Za-z0-9_.-]+)\s*\}\}",f,v)
    if isinstance(v,dict):return {k:templ(x,p) for k,x in v.items()}
    if isinstance(v,list):return [templ(x,p) for x in v]
    return v

def valid_url(url:str):
    p=urlparse(url)
    if p.scheme not in {"http","https"} or not p.hostname:raise HTTPException(400,"스킬 URL은 http:// 또는 https://여야 합니다.")
    host=p.hostname.lower().rstrip(".")
    if host in {"localhost","localhost.localdomain"}:raise HTTPException(400,"localhost는 스킬에서 사용할 수 없습니다.")
    try:
        for info in socket.getaddrinfo(host,p.port or (443 if p.scheme=="https" else 80),type=socket.SOCK_STREAM):
            ip=ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified:raise HTTPException(400,"공개 인터넷 주소만 스킬 URL로 등록할 수 있습니다.")
    except HTTPException:raise
    except Exception:pass

async def run_skill(s:dict,p:dict):
    url=templ(s["url"],p);valid_url(url);headers={}
    if s["headers"].strip():
        try:
            headers=templ(json.loads(s["headers"]),p)
            if not isinstance(headers,dict):raise ValueError
        except Exception:raise HTTPException(400,"Headers는 JSON 객체여야 합니다.")
    method=s["method"].upper();query=p if method in {"GET","DELETE"} else None;body=None
    if method not in {"GET","DELETE"} and s["body"].strip():
        raw=templ(s["body"],p)
        try:body=json.loads(raw)
        except Exception:body=raw
        if isinstance(body,(dict,list)):headers.setdefault("Content-Type","application/json")
    async with httpx.AsyncClient(timeout=15,follow_redirects=False,trust_env=False) as x:
        r=await x.request(method,url,headers=headers,params=query,json=body if isinstance(body,(dict,list)) else None,content=body if isinstance(body,str) else None)
    return {"status":r.status_code,"headers":{k:v for k,v in r.headers.items() if k.lower() in {"content-type","location","cache-control"}},"body":r.text[:12000]}

@app.get("/skills")
async def skills(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:return c.execute("SELECT id,name,description,url,method,headers,body,active,created_at,updated_at FROM mirae_skills WHERE user_id=%s ORDER BY created_at DESC",[u["id"]]).fetchall()

@app.post("/skills")
async def skill_create(data:SkillCreate,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    valid_url(str(data.url))
    if data.headers.strip():
        try:
            if not isinstance(json.loads(data.headers),dict):raise ValueError
        except Exception:raise HTTPException(400,"Headers는 JSON 객체여야 합니다.")
    with db() as c:
        try:
            r=c.execute("INSERT INTO mirae_skills(user_id,name,description,url,method,headers,body) VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id,name,description,url,method,headers,body,active,created_at,updated_at",[u["id"],data.name.strip(),data.description.strip(),str(data.url),data.method,data.headers,data.body]).fetchone();c.commit()
        except psycopg.errors.UniqueViolation:c.rollback();raise HTTPException(409,"같은 이름의 스킬이 이미 있습니다.")
    return r

@app.post("/skills/from-prompt")
async def skill_from_prompt(data:SkillPromptCreate,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    urls=re.findall(r"https?://[^\s<>{}\"]+",data.prompt)
    if not urls:raise HTTPException(400,"프롬프트에 API URL(http:// 또는 https://)을 포함해주세요.")
    parse_prompt=[
        {"role":"system","content":"Turn the user's natural-language skill description into ONLY one JSON object with keys name,description,url,method,headers,body. method must be GET,POST,PUT,PATCH,DELETE. Do not invent a URL: use exactly the URL from the user prompt. For GET/DELETE, put no body. For POST/PUT/PATCH, if the user names input fields, use {{field}} placeholders in JSON body. headers should be a JSON object string or empty. No Markdown."},
        {"role":"user","content":data.prompt}
    ]
    raw=await generate_once(parse_prompt,0.1,220)
    match=re.search(r"\{.*\}",raw,re.S)
    if not match:raise HTTPException(400,"스킬 설명을 구조화하지 못했습니다. URL과 사용 방법을 더 구체적으로 적어주세요.")
    try:obj=json.loads(match.group(0))
    except Exception:raise HTTPException(400,"스킬 설정을 읽지 못했습니다.")
    if str(obj.get("url","")) != urls[0].rstrip(".,)"):raise HTTPException(400,"AI가 입력한 API URL이 프롬프트의 URL과 일치하지 않습니다.")
    try:
        sc=SkillCreate(name=str(obj.get("name","api-skill")),description=str(obj.get("description","")),url=str(obj["url"]),method=str(obj.get("method","GET")).upper(),headers=str(obj.get("headers","")),body=str(obj.get("body","")))
    except Exception as e:raise HTTPException(400,f"생성된 스킬 형식이 올바르지 않습니다: {str(e)[:200]}")
    return await skill_create(sc,request)

@app.delete("/skills/{skill_id}")
async def skill_delete(skill_id:int,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:c.execute("DELETE FROM mirae_skills WHERE id=%s AND user_id=%s",[skill_id,u["id"]]);c.commit()
    return {"ok":True}

@app.post("/skills/{skill_id}/run")
async def skill_run(skill_id:int,data:SkillRun,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:s=c.execute("SELECT * FROM mirae_skills WHERE id=%s AND user_id=%s AND active=true",[skill_id,u["id"]]).fetchone()
    if not s:raise HTTPException(404,"스킬을 찾을 수 없습니다.")
    try:r=await run_skill(s,data.params)
    except HTTPException:raise
    except Exception as e:raise HTTPException(502,f"스킬 요청에 실패했습니다: {str(e)[:300]}")
    return {"skill":s["name"],"result":r}

PLUGIN_PERMISSIONS={"network","chat.read","profile.read","memory.read"}
def validate_plugin_permissions(items):
    clean=[]
    for x in items:
        if x not in PLUGIN_PERMISSIONS:raise HTTPException(400,f"지원하지 않는 플러그인 권한입니다: {x}")
        if x not in clean:clean.append(x)
    if "network" not in clean:raise HTTPException(400,"플러그인을 호출하려면 외부 네트워크 권한(network)을 허용해야 합니다.")
    return clean

async def invoke_plugin(plugin,user,request,input_data):
    permissions=validate_plugin_permissions(plugin["permissions"] or [])
    url=str(plugin["url"]);valid_url(url)
    payload={"input":input_data,"plugin":{"name":plugin["name"],"version":1},"permissions":permissions}
    if "chat.read" in permissions:
        payload["context"]={"history":[{"role":m["role"],"content":m["content"]} for m in input_data.get("_history",[]) if isinstance(m,dict) and m.get("role") in {"user","assistant"}][-12:]}
    else:payload["context"]={}
    if "profile.read" in permissions:
        with db() as c:
            p=c.execute("SELECT name,bio,birth_date FROM mirae_users WHERE id=%s",[user["id"]]).fetchone()
        payload["context"]["profile"]=dict(p or {})
    if "memory.read" in permissions:
        with db() as c:
            mm=c.execute("SELECT content FROM mirae_memories WHERE user_id=%s ORDER BY updated_at DESC LIMIT 20",[user["id"]]).fetchall()
        payload["context"]["memories"]=[x["content"] for x in mm]
    payload["input"].pop("_history",None)
    headers={"Content-Type":"application/json","X-Mirae-Plugin":plugin["name"]}
    async with httpx.AsyncClient(timeout=15,follow_redirects=False,trust_env=False) as x:
        if plugin["method"]=="GET":
            r=await x.get(url,params=payload["input"],headers={"X-Mirae-Plugin":plugin["name"]})
        else:
            r=await x.post(url,json=payload,headers=headers)
    return {"status":r.status_code,"headers":{k:v for k,v in r.headers.items() if k.lower() in {"content-type","location"}}, "body":r.text[:12000]}

@app.get("/plugins")
async def plugins(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:return c.execute("SELECT id,name,description,url,method,permissions,active,created_at,updated_at FROM mirae_plugins WHERE user_id=%s ORDER BY created_at DESC",[u["id"]]).fetchall()

@app.post("/plugins")
async def plugin_create(data:PluginCreate,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    perms=validate_plugin_permissions(data.permissions)
    valid_url(str(data.url))
    with db() as c:
        try:
            r=c.execute("INSERT INTO mirae_plugins(user_id,name,description,url,method,permissions) VALUES (%s,%s,%s,%s,%s,%s) RETURNING id,name,description,url,method,permissions,active,created_at,updated_at",
                        [u["id"],data.name.strip(),data.description.strip(),str(data.url),data.method,Jsonb(perms)]).fetchone()
            c.commit()
        except psycopg.errors.UniqueViolation:
            c.rollback();raise HTTPException(409,"같은 이름의 플러그인이 이미 있습니다.")
    return r

@app.put("/plugins/{plugin_id}")
async def plugin_update(plugin_id:int,data:PluginState,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        r=c.execute("UPDATE mirae_plugins SET active=%s,updated_at=now() WHERE id=%s AND user_id=%s RETURNING id,active",[data.active,plugin_id,u["id"]]).fetchone()
        if not r:raise HTTPException(404,"플러그인을 찾을 수 없습니다.")
        c.commit()
    return r

@app.delete("/plugins/{plugin_id}")
async def plugin_delete(plugin_id:int,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        r=c.execute("DELETE FROM mirae_plugins WHERE id=%s AND user_id=%s RETURNING id",[plugin_id,u["id"]]).fetchone();c.commit()
    return {"ok":True,"deleted":bool(r)}

@app.post("/plugins/{plugin_id}/invoke")
async def plugin_invoke(plugin_id:int,data:PluginInvoke,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:p=c.execute("SELECT * FROM mirae_plugins WHERE id=%s AND user_id=%s AND active=true",[plugin_id,u["id"]]).fetchone()
    if not p:raise HTTPException(404,"활성화된 플러그인을 찾을 수 없습니다.")
    try:r=await invoke_plugin(p,u,request,data.input)
    except HTTPException:raise
    except Exception as e:raise HTTPException(502,f"플러그인 요청에 실패했습니다: {type(e).__name__}: {str(e)[:300]}")
    return {"plugin":p["name"],"permissions":p["permissions"],"result":r}

async def run_plugin_command(message,request):
    m=re.match(r"^\s*/plugin\s+([^\s]+)(?:\s+(\{.*\}))?\s*$",message,re.S|re.I)
    if not m:return None
    u=session_user(request)
    if not u:raise HTTPException(401,"플러그인을 사용하려면 로그인해주세요.")
    try:pdata=json.loads(m.group(2) or "{}")
    except Exception:raise HTTPException(400,'플러그인 파라미터는 JSON이어야 합니다.')
    with db() as c:p=c.execute("SELECT * FROM mirae_plugins WHERE user_id=%s AND lower(name)=lower(%s) AND active=true",[u["id"],m.group(1)]).fetchone()
    if not p:raise HTTPException(404,f"'{m.group(1)}' 플러그인을 찾을 수 없습니다.")
    r=await invoke_plugin(p,u,request,pdata)
    return p,r

async def run_skill_command(message,request):
    m=re.match(r"^\s*/skill\s+([^\s]+)(?:\s+(\{.*\}))?\s*$",message,re.S|re.I)
    if not m:return None
    u=session_user(request)
    if not u:raise HTTPException(401,"스킬을 사용하려면 로그인해주세요.")
    try:p=json.loads(m.group(2) or "{}")
    except Exception:raise HTTPException(400,'파라미터는 JSON이어야 합니다. 예: /skill weather {"city":"천안"}')
    with db() as c:s=c.execute("SELECT * FROM mirae_skills WHERE user_id=%s AND lower(name)=lower(%s) AND active=true",[u["id"],m.group(1)]).fetchone()
    if not s:raise HTTPException(404,f"'{m.group(1)}' 스킬을 찾을 수 없습니다.")
    return s,await run_skill(s,p)

def model_headers():
    return {"Authorization":f"Bearer {MODEL_API_KEY}"} if MODEL_API_KEY else {}

def active_model_name():
    return MODEL_NAME

async def generate_once(msgs,temp,max_tokens):
    if not MODEL_API_URL:raise HTTPException(503,"MODEL_API_URL is not configured on the API server.")
    try:
        async with httpx.AsyncClient(timeout=180,trust_env=False) as x:
            r=await x.post(MODEL_API_URL,headers=model_headers(),json={"model":active_model_name(),"messages":msgs,"temperature":temp,"max_tokens":max_tokens,"stream":False})
    except Exception as e:
        raise HTTPException(502,f"AI 서버에 연결하지 못했습니다: {type(e).__name__}: {str(e)[:300]}")
    if r.status_code>=400:raise HTTPException(502,f"AI 서버 요청이 실패했습니다: {r.text[:500]}")
    try:
        content=r.json()["choices"][0]["message"]["content"]
        if isinstance(content,list):content="".join(str(x.get("text",x)) for x in content if isinstance(x,dict))
        content=str(content or "").strip()
        if not content:raise ValueError("empty content")
        return content
    except Exception:raise HTTPException(502,"AI 서버가 유효한 답변을 반환하지 않았습니다.")

async def generate_stream(msgs,temp,max_tokens):
    if not MODEL_API_URL:raise RuntimeError("MODEL_API_URL is not configured on the API server.")
    async with httpx.AsyncClient(timeout=180,trust_env=False) as x:
        async with x.stream("POST",MODEL_API_URL,headers=model_headers(),json={"model":active_model_name(),"messages":msgs,"temperature":temp,"max_tokens":max_tokens,"stream":True}) as r:
            if r.status_code>=400:raise RuntimeError((await r.aread()).decode(errors="ignore")[:500])
            content_type=(r.headers.get("content-type") or "").lower()
            if "text/event-stream" not in content_type:
                raw=await r.aread()
                try:o=json.loads(raw)
                except Exception as e:raise RuntimeError("AI ??? ???? ?? ??? ?? ? ????.") from e
                content=((o.get("choices") or [{}])[0].get("message") or {}).get("content","")
                if isinstance(content,list):content="".join(str(x.get("text",x)) for x in content if isinstance(x,dict))
                if content:yield str(content)
                return
            async for line in r.aiter_lines():
                if not line.startswith("data:"):continue
                raw=line[5:].strip()
                if raw=="[DONE]":break
                try:o=json.loads(raw)
                except Exception:continue
                ch=(o.get("choices") or [{}])[0].get("delta") or {}
                content=ch.get("content")
                if isinstance(content,str) and content:yield content

def event(name,data):return f"event: {name}\ndata: {json.dumps(data,ensure_ascii=False)}\n\n"

async def prepare(req,request,skip_search=False):
    u=session_user(request);sources=[];need=bool(req.web_search) and wants_web(req.message);skill_list=[];memories=[];profile_data={}
    if u:
        with db() as c:
            user_row=c.execute("SELECT name,bio,birth_date FROM mirae_users WHERE id=%s",[u["id"]]).fetchone()
            skill_list=c.execute("SELECT id,name,description FROM mirae_skills WHERE user_id=%s AND active=true ORDER BY name",[u["id"]]).fetchall()
            memories=c.execute("SELECT id,content FROM mirae_memories WHERE user_id=%s ORDER BY updated_at DESC LIMIT 20",[u["id"]]).fetchall()
            profile_data=user_row or {}
    if need and not skip_search:
        try:sources=await search_web(req.message)
        except Exception:sources=[]
    msgs=[{"role":"system","content":system_prompt(req,lang(req.message),sources,skill_list,memories,profile_data,req.attachments)}]
    msgs += [{"role":m.role,"content":m.content[:5000]} for m in req.history[-12:] if m.role in ("user","assistant") and m.content.strip()]
    msgs.append({"role":"user","content":req.message})
    return u,sources,msgs,need

@app.post("/chat")
async def chat(req:ChatRequest,request:Request):
    plugin=await run_plugin_command(req.message,request)
    if plugin:
        p,r=plugin;reply=f"[플러그인: {p['name']}]\nHTTP {r['status']}\n\n{r['body']}"
        u=session_user(request);cid=""
        if u:
            cid=save_chat(u["id"],req.message,reply,"plugin",[{"type":"plugin","name":p["name"]}],req.conversation_id or "")
            if current_title_missing(u["id"],cid):await finalize_conversation(u["id"],cid,req.message)
        return {"reply":reply,"model":"plugin","language":lang(req.message),"sources":[],"conversation_id":cid}
    sk=await run_skill_command(req.message,request);u=session_user(request)
    if sk:
        s,r=sk;reply=f"[스킬: {s['name']}]\nHTTP {r['status']}\n\n{r['body']}"
        if u:
            cid=save_chat(u["id"],req.message,reply,"skill",[{"type":"skill","name":s["name"]}],req.conversation_id or "")
            if current_title_missing(u["id"],cid):await finalize_conversation(u["id"],cid,req.message)
        else:cid=""
        return {"reply":reply,"model":"skill","language":lang(req.message),"sources":[]}
    u,sources,msgs,need=await prepare(req,request)
    image_intent=classify_image_intent(req.message)
    if image_intent["generate"]:
        image=await generate_image_payload(image_intent["prompt"]); reply="이미지를 생성했습니다."; cid=""
        if u:
            cid=save_chat(u["id"],req.message,reply,"image",[],req.conversation_id or "")
            if current_title_missing(u["id"],cid): await finalize_conversation(u["id"],cid,req.message)
        return {"reply":reply,"model":CLOUDFLARE_IMAGE_MODEL,"language":lang(req.message),"sources":sources,"conversation_id":cid,"image":image,"reasoning_summary":"요청의 의미를 분석해 이미지 생성 의도로 판단하고 이미지 프롬프트를 구성했습니다."}
    reply=await generate_once(msgs,req.temperature,req.max_tokens)
    cid=""
    if u:
        try:
            cid=save_chat(u["id"],req.message,reply,"web" if sources else "model",sources,req.conversation_id or "",[a.model_dump() for a in req.attachments])
        except Exception as e:
            print(f"CHAT_SAVE_ERROR: {type(e).__name__}: {str(e)[:500]}",flush=True)
            cid=req.conversation_id or ""
        if cid:
            try:
                if current_title_missing(u["id"],cid):
                    with db() as c:
                        c.execute("UPDATE mirae_conversations SET title=%s,updated_at=now() WHERE id=%s AND user_id=%s",
                                  [make_conversation_title(req.message),cid,u["id"]]);c.commit()
                    asyncio.create_task(finalize_conversation(u["id"],cid,req.message))
            except Exception as e:
                print(f"CHAT_TITLE_ERROR: {type(e).__name__}: {str(e)[:500]}",flush=True)
        try: asyncio.create_task(asyncio.to_thread(auto_register_memory,u["id"],req.message))
        except Exception: pass
    return {"reply":reply,"model":MODEL_NAME,"language":lang(req.message),"sources":sources,"conversation_id":cid,
            "reasoning_summary":("웹 검색 결과를 확인한 뒤 답변을 구성했습니다." if sources else "질문의 핵심을 파악하고 필요한 맥락을 반영해 답변을 구성했습니다.")}

@app.post("/chat/stream")
async def chat_stream(req:ChatRequest,request:Request):
    async def gen():
        try:
            sk=await run_skill_command(req.message,request);u=session_user(request)
            if sk:
                s,r=sk
                yield event("stage",{"id":"skill","label":f"'{s['name']}' 스킬 실행 중"})
                reply=f"[스킬: {s['name']}]\nHTTP {r['status']}\n\n{r['body']}"
                yield event("delta",{"text":reply})
                cid=""
                if u:
                    cid=save_chat(u["id"],req.message,reply,"skill",[{"type":"skill","name":s["name"]}],req.conversation_id or "")
                    if current_title_missing(u["id"],cid):
                        await finalize_conversation(u["id"],cid,req.message)
                yield event("conversation",{"id":cid,"title":"스킬 실행"})
                yield event("done",{"model":"skill","sources":[],"conversation_id":cid})
                return
            yield event("stage",{"id":"analyze","label":"질문 분석 중"})
            u,sources,msgs,need=await prepare(req,request,skip_search=True)
            if need:
                query=clean_query(req.message)
                yield event("stage",{"id":"search","label":"웹 검색 중"})
                yield event("search_query",{"query":query})
                try:sources=await search_web(req.message)
                except Exception:sources=[]
                msgs=[{"role":"system","content":system_prompt(req,lang(req.message),sources,[],[],{},req.attachments)}]
                if u:
                    with db() as c:
                        skill_list=c.execute("SELECT id,name,description FROM mirae_skills WHERE user_id=%s AND active=true ORDER BY name",[u["id"]]).fetchall()
                        memories=c.execute("SELECT id,content FROM mirae_memories WHERE user_id=%s ORDER BY updated_at DESC LIMIT 20",[u["id"]]).fetchall()
                        profile_data=c.execute("SELECT name,bio,birth_date FROM mirae_users WHERE id=%s",[u["id"]]).fetchone() or {}
                    msgs=[{"role":"system","content":system_prompt(req,lang(req.message),sources,skill_list,memories,profile_data,req.attachments)}]
                msgs += [{"role":m.role,"content":m.content[:5000]} for m in req.history[-12:] if m.role in ("user","assistant") and m.content.strip()]
                msgs.append({"role":"user","content":req.message})
                yield event("sources",{"sources":sources})
            yield event("stage",{"id":"generate","label":"답변 생성 중"});chunks=[]
            try:
                async for piece in generate_stream(msgs,req.temperature,req.max_tokens):chunks.append(piece);yield event("delta",{"text":piece})
            except Exception:
                if chunks:raise
                reply=await generate_once(msgs,req.temperature,req.max_tokens);chunks=[reply];yield event("delta",{"text":reply})
            reply="".join(chunks).strip()
            cid=""
            if u:
                cid=save_chat(u["id"],req.message,reply,"web" if sources else "model",sources,req.conversation_id or "",[a.model_dump() for a in req.attachments])
                is_first=(current_title_missing(u["id"],cid))
                if is_first:
                    yield event("stage",{"id":"title","label":"대화 제목 정리 중"})
                    await finalize_conversation(u["id"],cid,req.message)
                    with db() as c:title_row=c.execute("SELECT title FROM mirae_conversations WHERE id=%s",[cid]).fetchone()
                    yield event("conversation",{"id":cid,"title":title_row["title"] if title_row else "새 대화"})
            if u:asyncio.create_task(asyncio.to_thread(auto_register_memory,u["id"],req.message))
            yield event("done",{"model":MODEL_NAME,"sources":sources,"conversation_id":cid})
        except HTTPException as e:yield event("error",{"message":e.detail})
        except Exception as e:yield event("error",{"message":str(e)[:500]})
    return StreamingResponse(gen(),media_type="text/event-stream",headers={"Cache-Control":"no-cache, no-transform","X-Accel-Buffering":"no","Connection":"keep-alive"})

@app.get("/api-usage")
async def api_usage(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        summary=c.execute("""SELECT COUNT(*)::int AS requests,COALESCE(SUM(prompt_tokens),0)::int AS prompt_tokens,
            COALESCE(SUM(completion_tokens),0)::int AS completion_tokens,COALESCE(AVG(latency_ms),0)::int AS avg_latency_ms
            FROM mirae_api_usage WHERE user_id=%s""",[u["id"]]).fetchone()
        daily=c.execute("""SELECT date_trunc('day',created_at) AS day,COUNT(*)::int AS requests,
            COALESCE(SUM(prompt_tokens+completion_tokens),0)::int AS tokens
            FROM mirae_api_usage WHERE user_id=%s AND created_at>=now()-interval '30 days'
            GROUP BY 1 ORDER BY 1 DESC""",[u["id"]]).fetchall()
        logs=c.execute("""SELECT path,status,latency_ms,prompt_tokens,completion_tokens,created_at
            FROM mirae_api_usage WHERE user_id=%s ORDER BY created_at DESC LIMIT 100""",[u["id"]]).fetchall()
    return {"summary":summary,"daily":daily,"logs":logs}

@app.get("/webhooks")
async def list_webhooks(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:return c.execute("SELECT id,name,url,events,active,created_at FROM mirae_webhooks WHERE user_id=%s ORDER BY created_at DESC",[u["id"]]).fetchall()

@app.post("/webhooks")
async def create_webhook(data:WebhookCreate,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    if not data.url.startswith(("https://","http://")):raise HTTPException(400,"Webhook URL은 http 또는 https여야 합니다.")
    with db() as c:
        r=c.execute("INSERT INTO mirae_webhooks(user_id,name,url,events) VALUES (%s,%s,%s,%s) RETURNING id,name,url,events,active",[u["id"],data.name,data.url,json.dumps(data.events)]).fetchone();c.commit()
    return r

@app.put("/webhooks/{webhook_id}")
async def set_webhook(webhook_id:int,data:WebhookState,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        r=c.execute("UPDATE mirae_webhooks SET active=%s WHERE id=%s AND user_id=%s RETURNING id,active",[data.active,webhook_id,u["id"]]).fetchone();c.commit()
    if not r:raise HTTPException(404,"Webhook을 찾을 수 없습니다.")
    return r

@app.delete("/webhooks/{webhook_id}")
async def delete_webhook(webhook_id:int,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:
        r=c.execute("DELETE FROM mirae_webhooks WHERE id=%s AND user_id=%s RETURNING id",[webhook_id,u["id"]]).fetchone();c.commit()
    return {"ok":bool(r)}

async def dispatch_webhook(user_id,event,payload):
    try:
        with db() as c: hooks=c.execute("SELECT url FROM mirae_webhooks WHERE user_id=%s AND active=true AND events ? %s",[user_id,event]).fetchall()
        async with httpx.AsyncClient(timeout=5) as x:
            for h in hooks:
                try: await x.post(h["url"],json={"event":event,"created":int(time.time()),"data":payload})
                except Exception: pass
    except Exception: pass

@app.get("/api-keys")
async def list_keys(request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:return c.execute("SELECT id,name,key_prefix,state,last_used_at,revoked_at,created_at FROM mirae_api_keys WHERE user_id=%s ORDER BY created_at DESC",[u["id"]]).fetchall()

@app.post("/api-keys")
async def create_key(data:KeyCreate,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    raw="mk_"+secrets.token_urlsafe(32);prefix=raw[:10]
    with db() as c:c.execute("INSERT INTO mirae_api_keys(user_id,name,key_prefix,key_hash) VALUES (%s,%s,%s,%s)",[u["id"],data.name,prefix,digest(raw)]);c.commit()
    return {"key":raw,"prefix":prefix,"warning":"이 값은 지금 한 번만 표시됩니다."}

@app.delete("/api-keys/{key_id}")
async def revoke_key(key_id:int,request:Request):
    u=session_user(request)
    if not u:raise HTTPException(401,"로그인이 필요합니다.")
    with db() as c:c.execute("UPDATE mirae_api_keys SET state='revoked',revoked_at=now() WHERE id=%s AND user_id=%s",[key_id,u["id"]]);c.commit()
    return {"ok":True}

@app.post("/v1/chat/completions")
async def openai_chat(req:dict[str,Any],request:Request,authorization:str|None=Header(default=None)):
    started=time.perf_counter()
    key=authorization[7:] if authorization and authorization.startswith("Bearer ") else ""
    with db() as c:k=c.execute("SELECT * FROM mirae_api_keys WHERE key_hash=%s AND state='active'",[digest(key)]).fetchone() if key else None
    if not k:raise HTTPException(401,"Invalid or missing API key.")
    msgs=req.get("messages") or [];last=next((m.get("content","") for m in reversed(msgs) if m.get("role")=="user"),"");sources=[]
    try:
        if wants_web(last):sources=await search_web(last)
    except Exception:pass
    prompt=[{"role":"system","content":system_prompt(ChatRequest(message=last),lang(last),sources,[])}]+[m for m in msgs if m.get("role") in ("system","user","assistant")]
    reply=await generate_once(prompt,float(req.get("temperature",.7)),min(int(req.get("max_tokens",2200)),3200))
    latency=int((time.perf_counter()-started)*1000)
    with db() as c:
        c.execute("UPDATE mirae_api_keys SET last_used_at=now() WHERE id=%s",[k["id"]])
        c.execute("INSERT INTO mirae_api_usage(key_id,user_id,path,status,latency_ms) VALUES (%s,%s,%s,%s,%s)",[k["id"],k["user_id"],"/v1/chat/completions",200,latency]);c.commit()
    await dispatch_webhook(k["user_id"],"api.request",{"path":"/v1/chat/completions","model":req.get("model","mirae-free"),"latency_ms":latency})
    return {"id":"mirae-chat","object":"chat.completion","created":int(time.time()),"model":req.get("model","mirae-free"),"choices":[{"index":0,"message":{"role":"assistant","content":reply},"finish_reason":"stop"}],"usage":{"prompt_tokens":0,"completion_tokens":0,"total_tokens":0},"sources":sources}

@app.get("/v1/models")
async def models(authorization:str|None=Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "):raise HTTPException(401,"API key required.")
    with db() as c:k=c.execute("SELECT id FROM mirae_api_keys WHERE key_hash=%s AND state='active'",[digest(authorization[7:])]).fetchone()
    if not k:raise HTTPException(401,"Invalid API key.")
    return {"object":"list","data":[{"id":MODEL_NAME,"object":"model","owned_by":"mirae"}]}
